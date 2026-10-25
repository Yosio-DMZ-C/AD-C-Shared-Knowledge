#!/usr/bin/env python3
"""Public AD&C Draft smoke: editable PPTX -> real rendered PNG -> silent MP4.

No private input, network calls, brand assets, audio, paid services or publish.
Requires python-pptx, PyMuPDF, LibreOffice, FFmpeg and ffprobe.
Usage: python skills/video-production-setup/ppt_video_smoke.py --output smoke-output
"""
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SLIDES = [
    ("PPT動画のテスト", "編集可能なPowerPointから動画を作る経路を検査"),
    ("素材と構成", "この例は公開用の架空素材のみ。正式ロゴ・本人音声は含まない"),
    ("完成前の検査", "PPTX → 実レンダリング → 無音MP4。内容・音声の品質は別検査"),
]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def command(args, **kwargs):
    return subprocess.run(args, check=True, text=True, capture_output=True, **kwargs)


def make_pptx(path):
    prs = Presentation()
    prs.slide_width = Inches(13.3333)
    prs.slide_height = Inches(7.5)
    for index, (heading, subtitle) in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(250, 251, 254)
        stripe = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(.18), Inches(7.5))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor(17, 112, 124)
        stripe.line.fill.background()
        for content, top, size, color, bold in (
            (heading, 1.4, 34, RGBColor(17, 24, 39), True),
            (subtitle, 3.1, 20, RGBColor(32, 72, 85), False),
            (f"SMOKE / NOT A PRODUCTION VIDEO · {index:02d}/{len(SLIDES):02d}",
             6.8, 11, RGBColor(111, 122, 133), False),
        ):
            box = slide.shapes.add_textbox(Inches(.9), Inches(top), Inches(11.7), Inches(1.8))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = content
            run.font.name = "Noto Sans CJK JP"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    prs.save(path)
    with ZipFile(path) as z:
        parts = [x for x in z.namelist() if re.fullmatch(r"ppt/slides/slide[0-9]+\.xml", x)]
        if len(parts) != len(SLIDES):
            raise RuntimeError("saved PPTX slide count mismatch")
        if not any("素材".encode("utf-8") in z.read(x) for x in parts):
            raise RuntimeError("editable Japanese slide text missing")


def render(ppt, destination):
    import fitz
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office:
        raise RuntimeError("LibreOffice executable not found")
    with tempfile.TemporaryDirectory(prefix="adnc-lo-") as profile:
        command([office, "-env:UserInstallation=file://" + profile, "--headless",
                 "--convert-to", "pdf", "--outdir", str(destination), str(ppt)], timeout=120)
    pdf = destination / (ppt.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("LibreOffice did not produce a PDF")
    doc = fitz.open(str(pdf))
    if len(doc) != len(SLIDES):
        raise RuntimeError("PDF page count mismatch")
    images = []
    for i, page in enumerate(doc):
        text = re.sub(r"\s+", "", page.get_text("text"))
        if re.sub(r"\s+", "", SLIDES[i][0]) not in text:
            raise RuntimeError(f"slide {i+1} title missing from rendered output")
        png = destination / f"slide-{i+1:02d}.png"
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(str(png))
        images.append(png)
    doc.close()
    return pdf, images


def encode(images, output, folder):
    if not shutil.which("ffmpeg") or not shutil.which("ffprobe"):
        raise RuntimeError("FFmpeg or ffprobe executable not found")
    playlist = folder / "slides.ffconcat"
    lines = ["ffconcat version 1.0"]
    for image in images:
        lines += [f"file '{image.name}'", "duration 2.0"]
    lines.append(f"file '{images[-1].name}'")
    playlist.write_text("\n".join(lines) + "\n", encoding="utf-8")
    command(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
             "-safe", "0", "-f", "concat", "-i", str(playlist),
             "-vf", "fps=24,scale=trunc(iw/2)*2:trunc(ih/2)*2,format=yuv420p",
             "-c:v", "libx264", "-preset", "fast", "-t", "6.0",
             "-movflags", "+faststart", str(output)], cwd=str(folder), timeout=180)
    meta = json.loads(command(["ffprobe", "-v", "error", "-show_streams",
                               "-show_format", "-of", "json", str(output)]).stdout)
    streams = meta.get("streams", [])
    videos = [s for s in streams if s.get("codec_type") == "video"]
    audio = [s for s in streams if s.get("codec_type") == "audio"]
    duration = float(meta["format"]["duration"])
    if len(videos) != 1 or audio or not (5.5 <= duration <= 6.5):
        raise RuntimeError("MP4 format/track/duration invalid")
    return {"duration_seconds": duration, "width": videos[0]["width"],
            "height": videos[0]["height"], "audio_tracks": len(audio)}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--output", type=Path, default=Path("smoke-output"))
    args = ap.parse_args()
    out = args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    images_dir = out / "rendered"
    images_dir.mkdir(exist_ok=True)
    ppt = out / "ADNC_PUBLIC_SMOKE_EDITABLE.pptx"
    mp4 = out / "ADNC_PUBLIC_SMOKE_SILENT.mp4"
    result = {"kind": "public_synthetic_video_pipeline_smoke", "production_ready": False,
              "privacy": "no private input files or service credentials", "checks": {}}
    try:
        make_pptx(ppt)
        result["checks"]["editable_pptx"] = "PASS"
        pdf, images = render(ppt, images_dir)
        result["checks"]["pptx_rendered_to_images"] = "PASS"
        result["media"] = encode(images, mp4, images_dir)
        result["checks"]["silent_mp4_encoding"] = "PASS"
        result["checks"].update({k: "NOT_RUN" for k in (
            "ai_voice", "sync_to_real_voice", "official_brand_assets",
            "human_visual_review", "another_ai_cold_start", "friend_environment")})
        result["files"] = [
            {"name": p.relative_to(out).as_posix(), "bytes": p.stat().st_size,
             "sha256": digest(p)} for p in [ppt, pdf, *images, mp4]]
        result["overall"] = "PASS_FOR_SYNTHETIC_PIPELINE_ONLY"
    except Exception as exc:
        result["overall"] = "FAIL"
        result["error"] = str(exc)
    report = out / "smoke-report.json"
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n",
                      encoding="utf-8")
    print(json.dumps({"overall": result["overall"], "report": str(report),
                      "checks": result["checks"]}, ensure_ascii=False))
    return 0 if result["overall"].startswith("PASS") else 1


if __name__ == "__main__":
    sys.exit(main())
